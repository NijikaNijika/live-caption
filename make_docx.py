#!/usr/bin/env python3
"""make_docx.py — 把一次同传会话的双语记录导出成 Word 文档

由 captioner.mjs 在退出时调用（不是手工运行）：
    python3 make_docx.py <payload.json> <输出目录>

payload.json 结构见 captioner.mjs 里的 buildDocPayload()。
设计目标是"离开程序就能拿到一份能直接读/打印/交作业的文档"。
"""

import json
import os
import sys
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

# 中文字形必须显式指定 eastAsia 字体，只设 run.font.name 对中文不生效
LATIN_FONT = "Helvetica Neue"
CJK_FONT = "PingFang SC"


def style_run(run, size=None, bold=None, italic=None, color=None, cjk=None, latin=None):
    """统一设置字体；同时写入 w:eastAsia，保证中文不回退成方框。"""
    latin = latin or LATIN_FONT
    cjk = cjk or CJK_FONT
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_para(doc, text, size=11, bold=False, italic=False, color=None,
             space_after=6, space_before=0, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    style_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading_cjk(doc, text, level):
    """add_heading 建的是原生 heading，字号随样式走但字体是默认的；
    中文标题必须补 w:eastAsia，否则可能回退成方框。"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:ascii"), LATIN_FONT)
        rFonts.set(qn("w:hAnsi"), LATIN_FONT)
        rFonts.set(qn("w:eastAsia"), CJK_FONT)
    return h


def add_meta_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for k, v in rows:
        cells = table.add_row().cells
        style_run(cells[0].paragraphs[0].add_run(k), size=10, bold=True)
        style_run(cells[1].paragraphs[0].add_run(str(v)), size=10)
    return table


def fmt_clock(ms_from_start):
    """把毫秒偏移格式化成 HH:MM:SS，读起来比绝对时间戳有用。"""
    total = max(0, int(ms_from_start // 1000))
    h, rem = divmod(total, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def main():
    if len(sys.argv) < 3:
        print("用法: make_docx.py <payload.json> <输出目录>", file=sys.stderr)
        return 2

    payload_path, out_dir = sys.argv[1], sys.argv[2]
    with open(payload_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    segments = data.get("segments") or []
    if not segments:
        print("没有可导出的内容（会话里没有任何定稿段落），跳过生成。", file=sys.stderr)
        return 3

    os.makedirs(out_dir, exist_ok=True)

    doc = Document()

    # ── 标题：当日日期 + 时间 + 城市 ──
    title = add_heading_cjk(doc, data["title"], 0)
    for run in title.runs:
        style_run(run, size=26, bold=True)

    for p in doc.paragraphs[:1]:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # ── 概览 ──
    add_heading_cjk(doc, "本次记录", 1)
    stats = data.get("stats") or {}
    add_meta_table(doc, [
        ("日期", data.get("date", "")),
        ("时间", data.get("time", "")),
        ("城市", data.get("city", "")),
        ("段落数", stats.get("count", len(segments))),
        ("记录时长", stats.get("durationText", "")),
        ("上屏延迟中位", stats.get("medianLatencyText", "")),
        ("语音识别", stats.get("asrModel", "")),
        ("翻译模型", stats.get("mtModel", "")),
    ])

    add_para(doc, "", size=6, space_after=2)
    add_para(doc,
             "说明：英文为现场语音识别原文（可能存在识别误差），中文为实时机器翻译结果，"
             "供听讲时理解使用，未经人工校对。",
             size=9, italic=True, color=(0x66, 0x66, 0x66))

    # ── 双语对照正文 ──
    add_heading_cjk(doc, "双语对照", 1)
    base = data.get("startedAt") or segments[0].get("t") or 0
    for i, seg in enumerate(segments, 1):
        offset = (seg.get("t") or base) - base
        head = f"[{i:03d}]  {fmt_clock(offset)}"
        if seg.get("latencyMs"):
            head += f"   上屏 {int(seg['latencyMs'])} ms"
        add_para(doc, head, size=8.5, color=(0x88, 0x88, 0x88),
                 space_before=8 if i > 1 else 2, space_after=2)

        add_para(doc, seg.get("en", ""), size=10.5, italic=True,
                 color=(0x33, 0x33, 0x33), space_after=2)
        add_para(doc, seg.get("zh", ""), size=12.5, bold=False,
                 color=(0x0B, 0x3C, 0x5D), space_after=4)

    # 文件名：日期 + 时间 + 城市（不含冒号，避免 Finder 把冒号显示成斜杠）
    filename = data.get("filename") or f"{data.get('date')} {data.get('timeCompact', '')} {data.get('city')}.docx"
    out_path = os.path.join(out_dir, filename)
    doc.save(out_path)
    print(out_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
